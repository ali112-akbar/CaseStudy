from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
out = root / 'deliverables'
out.mkdir(exist_ok=True)
doc = Document()
s = doc.sections[0]
s.top_margin = s.bottom_margin = Inches(.7)
s.left_margin = s.right_margin = Inches(.8)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Heading 1'].font.size = Pt(14)
doc.styles['Heading 1'].font.color.rgb = RGBColor.from_string('6D28D9')
doc.core_properties.title = 'Household Income and Political Party Identification'
doc.core_properties.author = 'Ali Akbar and team'
doc.add_paragraph('Household Income and Political Party Identification', 'Title')
doc.add_paragraph('Presentation script | EMBES 2010 | Group 1', 'Subtitle')
doc.add_paragraph('Our research question is: Does annual household income play a role in ethnic minorities’ political party identification? Our current analysis suggests a small association between income and party identification, with some important limitations.')
doc.add_heading('The survey and our approach', 1)
doc.add_paragraph('We used the 2010 Ethnic Minority British Election Study, covering five ethnic minority groups. We studied party identification—the party someone identifies with—which is different from their actual vote.')
doc.add_paragraph('We first investigated missing responses. Of 2,787 respondents, 1,597 had valid answers for both income and party identification. We excluded refused and unknown answers from the main comparison, but kept “None/No” because having no party identification is a meaningful response.')
doc.add_paragraph('Missing income was uneven across groups. For example, it affected 49.2% of Pakistani respondents compared with 32.2% of Black Caribbean respondents. This means excluding incomplete responses could influence our findings.')
doc.add_heading('What the income groups show', 1)
doc.add_paragraph('We then compared survey-weighted party percentages across four household-income groups. Labour remained the largest identification category in every income group. Labour identification was around 69–71% in the two lower-income groups, compared with 61.2% above £40,000.')
doc.add_paragraph('Conservative identification showed the opposite broad pattern: around 6–8% in the lower-income groups, compared with 16.9% above £40,000. The pattern was not a steady increase at every income step, but Conservative identification was more common in the upper two groups.')
doc.add_paragraph('Liberal Democrat identification and having no party identification showed relatively little variation across the income groups.')
doc.add_page_break()
doc.add_heading('How strong the association is', 1)
doc.add_paragraph('Our exploratory chi-square test gave χ² = 26.43, with 12 degrees of freedom and p = 0.0093. Under the ordinary test’s assumptions, this provides evidence against income and party identification being independent.')
doc.add_paragraph('However, Cramér’s V was 0.074, indicating a small association. Statistical significance does not mean that income strongly explains party identification.')
doc.add_paragraph('We checked an alternative arrangement using three income groups. The p-value remained below 0.05 and the effect size remained small. This suggests our broad finding is not limited to the original four-group arrangement.')
doc.add_heading('What our team comparison adds', 1)
doc.add_paragraph('Our team’s additional chart examined the relationship from another direction: the income distribution within Labour and Conservative identifiers. Among those with valid income, 69.1% of Conservative identifiers reported household income above £20,000, compared with 48.5% of Labour identifiers, using survey weights. This complements the main chart using the same data.')
doc.add_heading('Our conclusion and its limits', 1)
doc.add_paragraph('Our conclusion is therefore that income is associated with party identification in this analysis, but the relationship is small. Labour remains dominant across income groups, while Conservative identification is more common in higher-income groups.')
doc.add_paragraph('We cannot conclude that income causes these differences. Our descriptive charts use survey weights, but the current chi-square tests do not account for the full survey design. We also have substantial missing income and have not adjusted for characteristics such as ethnicity, age or education. These findings should therefore be presented as exploratory, with survey-adjusted testing as the next step.')
doc.save(out / 'Income_Party_Presentation.docx')
print(out / 'Income_Party_Presentation.docx')
